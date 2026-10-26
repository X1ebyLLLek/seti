"""Презентация к защите (.pptx, 12 слайдов, 16:9) из сохранённых результатов.

    python docs/report/build_slides.py      ->  docs/report/Презентация.pptx

Числа и диаграммы берутся из тех же файлов, что и в записке (build_report.py).
Текст доклада (7–10 минут) — в заметках докладчика к каждому слайду.
"""
from __future__ import annotations

import collections
import glob
import json
import os
import sys
from math import comb

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_report as R  # noqa: E402  (данные и пути — общие с запиской)

OUT = os.path.join(HERE, "Презентация.pptx")

INK = RGBColor(0x12, 0x26, 0x3A)        # доминирующий тёмный
TEAL = RGBColor(0x1B, 0x99, 0x8B)       # вторичный: «защищённый канал»
AMBER = RGBColor(0xF4, 0xA2, 0x59)      # акцент: отказ, внимание
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xEE, 0xF3, 0xF6)       # подложка карточек
MUTED = RGBColor(0x5B, 0x67, 0x70)
LINE = RGBColor(0xC9, 0xD3, 0xDA)
HEAD, BODY = "Cambria", "Calibri"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)


# ====================================================================== примитивы
def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None, space_after=0):
    """runs: строка, список строк (абзацы) или список списков (абзац из фрагментов (текст, {опции}))."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        frags = [(para, {})] if isinstance(para, str) else para
        for frag, opt in frags:
            r = p.add_run()
            r.text = frag
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = opt.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=MIST, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None, radius=0.08):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def node(slide, x, y, d, label, fill=TEAL, color=PAPER, size=14):
    """Мотив презентации: узел сети — круг с меткой."""
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = BODY, Pt(size), True, color
    return c


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    ln.line.color.rgb = color
    ln.line.width = Pt(width)
    ln.line._get_or_add_ln().append(_tail_arrow())
    return ln


def _tail_arrow():
    from pptx.oxml.ns import qn
    from lxml import etree
    el = etree.SubElement(etree.Element("dummy"), qn("a:tailEnd"))
    el.set("type", "triangle")
    el.set("w", "med")
    el.set("len", "med")
    return el


def title(slide, t, sub=None):
    text(slide, M, Inches(0.45), W - 2 * M, Inches(0.8), t, size=34, bold=True, font=HEAD, name="Заголовок")
    if sub:
        text(slide, M, Inches(1.18), W - 2 * M, Inches(0.45), sub, size=16, color=MUTED, name="Подзаголовок")


def footer(slide, n, dark=False):
    c = PAPER if dark else MUTED
    text(slide, M, H - Inches(0.45), Inches(8), Inches(0.3),
         "Вихарев Д.А., ИСБ-123 · Защищённая оверлейная P2P-сеть", size=10, color=c)
    text(slide, W - M - Inches(1), H - Inches(0.45), Inches(1), Inches(0.3), str(n), size=10, color=c,
         align=PP_ALIGN.RIGHT)


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def stat(slide, x, y, w, value, label, color=INK, vsize=40):
    text(slide, x, y, w, Inches(0.75), value, size=vsize, bold=True, font=HEAD, color=color)
    text(slide, x, y + Inches(0.08 + vsize / 72), w, Inches(0.6), label, size=13, color=MUTED)


def style_chart(chart, colors):
    chart.font.name = BODY
    chart.font.size = Pt(12)
    chart.font.color.rgb = MUTED
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    for s, col in zip(plot.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = col
    va = chart.value_axis
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    ca = chart.category_axis
    ca.format.line.color.rgb = LINE
    ca.tick_labels.font.size = Pt(11)
    chart.has_legend = len(plot.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)


# ====================================================================== данные
def bucket_fill(scheme, ranks=6):
    acc = collections.defaultdict(list)
    for f in glob.glob(os.path.join(R.SERIES, scheme, "run-*", "tables.json")):
        for v in json.load(open(f, encoding="utf-8")).values():
            b = {int(k): c for k, c in v["buckets"].items()}
            for j in range(ranks):
                acc[j].append(b.get(255 - j, 0))
    return [sum(acc[j]) / len(acc[j]) for j in range(ranks)]


def bucket_forecast(n_other=14, k=3, ranks=6):
    out = []
    for j in range(ranks):
        p = 2 ** -(j + 1)
        out.append(sum(min(k, x) * comb(n_other, x) * p ** x * (1 - p) ** (n_other - x) for x in range(n_other + 1)))
    return out


def rpc_hist(scheme):
    c = collections.Counter(int(x["rpc_count"]) for x in R.series_rows(scheme, "lookups.csv"))
    total = sum(c.values())
    return {k: 100 * v / total for k, v in c.items()}, total


# ====================================================================== слайды
def s_title(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, INK)
    pts = [(9.4, 1.0), (11.2, 1.7), (10.2, 3.0), (12.0, 3.6), (9.0, 4.4), (11.0, 5.2)]
    for a, b in [(0, 1), (0, 2), (1, 3), (2, 3), (2, 4), (4, 5), (3, 5)]:
        (x1, y1), (x2, y2) = pts[a], pts[b]
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1 + 0.275), Inches(y1 + 0.275),
                                    Inches(x2 + 0.275), Inches(y2 + 0.275))
        ln.line.color.rgb = RGBColor(0x3A, 0x55, 0x6E)
        ln.line.width = Pt(1.5)
    for i, (x, y) in enumerate(pts):
        node(s, Inches(x), Inches(y), Inches(0.55), "", fill=TEAL if i % 3 else AMBER)
    text(s, M, Inches(1.6), Inches(8.2), Inches(2.4),
         "Защищённая оверлейная сеть передачи данных на основе пиринговых протоколов",
         size=38, bold=True, font=HEAD, color=PAPER)
    text(s, M, Inches(4.25), Inches(8), Inches(0.9),
         "Kademlia DHT · TLS 1.3 со взаимной аутентификацией · туннели через ретрансляторы",
         size=18, color=RGBColor(0xB8, 0xD8, 0xD3))
    text(s, M, Inches(5.6), Inches(8), Inches(1.0),
         ["Курсовая работа · Вихарев Д.А., группа ИСБ-123", "ВлГУ, кафедра информатики и защиты информации, 2026"],
         size=15, color=PAPER)
    s.notes_slide.notes_text_frame.text = (
        "Здравствуйте. Тема работы — защищённая оверлейная сеть на основе пиринговых протоколов. "
        "Я реализовал узел такой сети: поиск узлов через Kademlia, защиту каналов TLS 1.3 и передачу сообщений "
        "через цепочку ретрансляторов с восстановлением после отказа, и проверил его экспериментально.")


def s_goal(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Цель и выбранный уровень",
          "Узел децентрализованной сети без центрального каталога — упрощённый уровень выполнения")
    items = [("Идентификация", "NodeID = SHA-256 от Ed25519-ключа, без центрального каталога"),
             ("Поиск", "Kademlia: k-bucket, итеративный lookup, подписанные записи с репликацией"),
             ("Защита", "TLS 1.3 mTLS на каждом звене и end-to-end поверх туннеля"),
             ("Отказоустойчивость", "автоматическое перестроение туннеля после отказа ретранслятора")]
    y = Inches(1.95)
    for i, (h, d) in enumerate(items):
        node(s, M, y, Inches(0.5), str(i + 1))
        text(s, M + Inches(0.7), y - Inches(0.02), Inches(6.2), Inches(0.35), h, size=18, bold=True)
        text(s, M + Inches(0.7), y + Inches(0.33), Inches(6.2), Inches(0.6), d, size=15, color=MUTED)
        y += Inches(1.12)
    box(s, Inches(7.6), Inches(1.95), Inches(5.13), Inches(4.4), fill=MIST)
    stats = [("15", "узлов N"), ("3", "k (ёмкость bucket)"), ("3", "ALPHA, R"), ("2", "ретранслятора")]
    for i, (v, l) in enumerate(stats):
        x = Inches(7.95) + (i % 2) * Inches(2.45)
        yy = Inches(2.25) + (i // 2) * Inches(1.55)
        stat(s, x, yy, Inches(2.2), v, l, color=TEAL if i % 2 else INK, vsize=44)
    text(s, Inches(7.95), Inches(5.4), Inches(4.5), Inches(0.8),
         "Продвинутому уровню нужны Docker/namespaces и netem — на рабочей машине (Windows) их нет",
         size=13, color=MUTED)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Цель — узел, который сам себя идентифицирует, находит других без каталога, защищает данные и переживает "
        "отказ промежуточного узла. Выбран упрощённый уровень: 15 узлов, ёмкость bucket 3, параллелизм поиска 3, "
        "три реплики, два ретранслятора. Продвинутый требует средств ограничения сети, которых на машине нет.")


def s_arch(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Архитектура узла", "Один процесс asyncio; подсистемы обращаются друг к другу только сверху вниз")
    layers = [("Интерфейс: консоль run / shell / ctl, управляющий порт с токеном", MIST, INK),
              ("Прикладной сервис: сообщения, E2E TLS, message_id, восстановление", MIST, INK),
              (None, None, None),
              ("Идентичность и защита: Ed25519, NodeID, TLS 1.3 mTLS, привязка к NodeID", MIST, INK),
              ("Транспорт: TCP, кадр 24 байта, MessagePack-схемы, RPC по request_id", INK, PAPER)]
    x, w, h = M, Inches(7.4), Inches(0.82)
    y = Inches(1.9)
    for label, fill, col in layers:
        if label is None:
            b1 = box(s, x, y, Inches(3.62), h, fill=TEAL)
            b2 = box(s, x + Inches(3.78), y, Inches(3.62), h, fill=TEAL)
            for b, t in ((b1, "Туннели: BUILD, DATA/ACK, состояния"), (b2, "DHT: k-bucket, lookup, записи")):
                tf = b.text_frame
                tf.margin_left = Inches(0.15)
                p = tf.paragraphs[0]
                r = p.add_run()
                r.text = t
                r.font.name, r.font.size, r.font.bold, r.font.color.rgb = BODY, Pt(15), True, PAPER
        else:
            b = box(s, x, y, w, h, fill=fill)
            tf = b.text_frame
            tf.margin_left = Inches(0.2)
            r = tf.paragraphs[0].add_run()
            r.text = label
            r.font.name, r.font.size, r.font.color.rgb = BODY, Pt(15), col
            r.font.bold = fill == INK
        y += h + Inches(0.12)
    notes = [("Фасад узла", "join, lookup_node, put_record, get_record, build_tunnel, send_via_tunnel — прикладной "
                            "слой не трогает k-bucket и хранилище"),
             ("Пул соединений", "одно TLS-рукопожатие на пару узлов; по соединению идут и RPC, и ячейки туннелей"),
             ("Процессы-узлы", "у каждого свой ключ, порт, таблица, хранилище и журнал (node.log + events.jsonl)")]
    yy = Inches(1.9)
    for h_, d in notes:
        text(s, Inches(8.4), yy, Inches(4.3), Inches(0.35), h_, size=17, bold=True)
        text(s, Inches(8.4), yy + Inches(0.38), Inches(4.3), Inches(1.0), d, size=14, color=MUTED)
        yy += Inches(1.45)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Узел разделён на подсистемы, как требует задание: транспорт, идентичность и защита, DHT, туннели, "
        "прикладной сервис и интерфейс. Прикладной слой работает только через фасад. Соединения переиспользуются: "
        "одно TLS-рукопожатие на пару узлов, по нему мультиплексируются запросы и ячейки туннелей.")


def s_frame(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Кадр и протокол", "Собственное кадрирование поверх TCP: граница сообщения задаётся длиной")
    fields = [("version", 1), ("type", 1), ("flags", 2), ("request_id", 16), ("payload_length", 4)]
    x0, y0, unit = M, Inches(2.0), Inches(0.43)
    x = x0
    for i, (name, size) in enumerate(fields):
        w = unit * max(size, 2) if size < 16 else unit * 9
        b = box(s, x, y0, w - Inches(0.05), Inches(0.9), fill=INK if i != 3 else TEAL, shape=MSO_SHAPE.RECTANGLE)
        tf = b.text_frame
        tf.margin_left = tf.margin_right = Inches(0.03)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = f"{name}\n{size} Б"
        r.font.name, r.font.size, r.font.color.rgb, r.font.bold = BODY, Pt(13), PAPER, True
        x += w
    pb = box(s, x, y0, W - M - x, Inches(0.9), fill=MIST, line=LINE, shape=MSO_SHAPE.RECTANGLE)
    tf = pb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "payload ≤ 65 536 Б\nMessagePack"
    r.font.name, r.font.size, r.font.color.rgb = BODY, Pt(13), INK
    text(s, M, Inches(3.05), Inches(6), Inches(0.4), "заголовок — 24 байта, сетевой порядок байтов", size=13,
         color=MUTED)
    cols = [("Защита декодера", ["длина проверяется до приёма payload", "частичные, склеенные и разрезанные "
                                 "кадры", "ошибка формата → ERROR и закрытие"]),
            ("21 тип сообщений", ["PING, FIND_NODE, STORE, FIND_VALUE", "SECURE_HELLO, TUNNEL_*", "APP_MESSAGE, "
                                  "APP_ACK, ERROR"]),
            ("Почему MessagePack", ["на 25–30 % меньше JSON с base64", "кодирование в 6–80 раз быстрее JSON",
                                    "строгие схемы, ext-типы запрещены"])]
    for i, (h_, items) in enumerate(cols):
        x = M + i * Inches(4.1)
        box(s, x, Inches(3.75), Inches(3.85), Inches(2.6), fill=MIST)
        text(s, x + Inches(0.25), Inches(3.95), Inches(3.4), Inches(0.4), h_, size=17, bold=True)
        text(s, x + Inches(0.25), Inches(4.5), Inches(3.4), Inches(1.8), ["– " + t for t in items], size=14,
             color=INK, space_after=6)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "TCP не сохраняет границы сообщений, поэтому каждый кадр начинается 24-байтовым заголовком с длиной. "
        "Длина проверяется сразу после заголовка, до выделения буфера, — завышенная длина отклоняется. "
        "Формат payload — MessagePack: по замеру он меньше JSON и быстрее и JSON, и CBOR.")


def s_kademlia(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Kademlia: DHT не вырождается при N = 15, k = 3",
          "Узел знает ≈ 9 контактов из 14 — прогноз и измерение совпадают")
    text(s, M, Inches(1.95), Inches(5.0), Inches(4.2), [
        [("XOR-метрика: ", {"bold": True}), ("расстояние = NodeID ⊕ NodeID, 256 bucket", {})],
        [("k-bucket: ", {"bold": True}), ("не более 3 контактов; живой давний контакт не вытесняется — сначала PING", {})],
        [("Lookup: ", {"bold": True}), ("опрос до 3 кандидатов из окна k ближайших, пока все они не ответят", {})],
        [("Прогноз: ", {"bold": True}), ("E[min(3, Bin(14, 2^−(j+1)))] по bucket — 8,97 контакта на узел", {})],
    ], size=15, space_after=12)
    pred = bucket_forecast()
    star, ring = bucket_fill("star"), bucket_fill("ring")
    cd = CategoryChartData()
    cd.categories = [f"{255 - j}" for j in range(6)]
    cd.add_series("прогноз", [round(v, 2) for v in pred])
    cd.add_series("star, измерено", [round(v, 2) for v in star])
    cd.add_series("ring, измерено", [round(v, 2) for v in ring])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.0), Inches(1.85), Inches(6.73), Inches(4.3), cd)
    ch = gf.chart
    style_chart(ch, [LINE, INK, TEAL])
    ch.has_title = True
    ch.chart_title.text_frame.text = "Среднее заполнение bucket (75 узлов на схему)"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.value_axis.minimum_scale = 0
    ch.value_axis.maximum_scale = 3.5
    ch.value_axis.major_unit = 0.5
    ch.value_axis.tick_labels.number_format = '0.0'
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.category_axis.has_title = True
    ch.category_axis.axis_title.text_frame.text = "номер bucket"
    ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    sizes = R.STATS[("star", "Контактов в таблице узла")], R.STATS[("ring", "Контактов в таблице узла")]
    text(s, M, Inches(5.55), Inches(5.2), Inches(0.7),
         f"Средний размер таблицы: {R.num(sizes[0]['mean'])} (star), {R.num(sizes[1]['mean'])} (ring); "
         f"максимум {max(int(float(x['max'])) for x in sizes)} при N − 1 = 14", size=14, color=MUTED)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Главный риск маленькой сети — вырождение: если каждый знает всех, поиск через DHT теряет смысл. "
        "Я заранее посчитал ожидаемое заполнение bucket: в среднем 8,97 контакта из 14. Измерения на 75 узлах каждой "
        "схемы дали около 9 — прогноз подтвердился, у всех узлов таблица меньше N − 1.")


def s_records(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Подписанные записи и репликация", "Адрес узла находится через FIND_VALUE, а не из каталога")
    chain = ["схема и размер", "NodeID = SHA-256(ключ)", "ключ хранения", "подпись Ed25519", "срок действия",
             "sequence_number ≥ максимума"]
    text(s, M, Inches(1.95), Inches(6), Inches(0.4), "Проверка перед сохранением и использованием", size=17,
         bold=True)
    y = Inches(2.5)
    for i, c in enumerate(chain):
        node(s, M, y, Inches(0.42), str(i + 1), size=12)
        text(s, M + Inches(0.6), y + Inches(0.06), Inches(5), Inches(0.35), c, size=15)
        y += Inches(0.55)
    box(s, Inches(7.0), Inches(1.95), Inches(5.73), Inches(2.0), fill=MIST)
    text(s, Inches(7.3), Inches(2.15), Inches(5.2), Inches(1.7), [
        [("STORE на R = 3 ближайших, ", {"bold": True}), ("успех — от 2 подтверждений", {})],
        [("Republish каждые 90 с ", {"bold": True}), ("(TTL 180 с) с новым номером", {})],
        [("Откат и подмена ", {"bold": True}), ("отклоняются: ROLLBACK, SIGNATURE", {})]], size=15, space_after=8)
    box(s, Inches(7.0), Inches(4.15), Inches(5.73), Inches(1.85), fill=PAPER, line=AMBER)
    text(s, Inches(7.3), Inches(4.3), Inches(5.2), Inches(2.0), [
        [("Найдено экспериментом", {"bold": True, "color": AMBER})],
        [("Узлы, вошедшие позже и ставшие ближе к ключу, не имели записи до republish. Исправлено передачей "
          "записи новому ближайшему соседу (§2.5 статьи Kademlia) и учётом своей копии.", {})]],
        size=14, space_after=6)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Каждый узел публикует подписанную запись со своим адресом. Хранитель проверяет её полностью: привязку "
        "ключа к NodeID, подпись, срок и номер версии — откат отклоняется. Первая экспериментальная серия нашла "
        "реальный дефект: запись не находилась, если ближе к ключу оказались узлы, вошедшие позже. Я добавил "
        "механизм из оригинальной статьи — передачу записи новому ближайшему соседу.")


def s_security(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Защита канала: TLS 1.3 с привязкой к NodeID",
          "Самоподписанный сертификат принимается, только если его ключ хэшируется в ожидаемый NodeID")
    steps = [("SECURE_HELLO", "NodeID, сертификат, nonce"), ("Проверка", "SHA-256(ключ) = NodeID"),
             ("TLS 1.3 mTLS", "доверие одному сертификату"), ("Транскрипт", "хэш HELLO внутри TLS")]
    x = M
    for i, (h_, d) in enumerate(steps):
        b = box(s, x, Inches(1.95), Inches(2.75), Inches(1.15), fill=INK if i == 2 else MIST)
        col = PAPER if i == 2 else INK
        text(s, x + Inches(0.2), Inches(2.08), Inches(2.4), Inches(0.4), h_, size=16, bold=True, color=col)
        text(s, x + Inches(0.2), Inches(2.5), Inches(2.4), Inches(0.5), d, size=13,
             color=PAPER if i == 2 else MUTED)
        if i < 3:
            arrow(s, x + Inches(2.78), Inches(2.52), x + Inches(3.05), Inches(2.52))
        x += Inches(3.07)
    # два вложенных канала: hop-by-hop и E2E
    y = Inches(3.9)
    labels = ["I", "R1", "R2", "D"]
    d = Inches(0.75)
    xs = [M + Inches(0.1) + i * Inches(2.25) for i in range(4)]
    for i in range(3):
        box(s, xs[i] + d / 2, y + Inches(0.25), Inches(2.25), Inches(0.25), fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
    box(s, xs[0] + d / 2, y + Inches(1.3), xs[3] - xs[0], Inches(0.25), fill=AMBER, shape=MSO_SHAPE.RECTANGLE)
    for x_, l in zip(xs, labels):
        node(s, x_, y, d, l, fill=INK, size=15)
        node(s, x_, y + Inches(1.05), d, l, fill=INK, size=15)
    text(s, M, y - Inches(0.45), Inches(7), Inches(0.35), "Каждое звено — свой TLS 1.3", size=14, bold=True, color=TEAL)
    text(s, M, y + Inches(1.9), Inches(7.2), Inches(0.7), [
        [("E2E TLS 1.3 отправитель → получатель поверх туннеля: ", {"bold": True, "color": AMBER}),
         ("ретранслятор видит только шифртекст", {"color": MUTED})]], size=14)
    text(s, Inches(8.6), Inches(3.55), Inches(4.1), Inches(2.9), [
        [("TLS 1.2 и 0-RTT ", {"bold": True}), ("отключены, билетов нет", {})],
        [("Отрицательные тесты: ", {"bold": True}), ("подмена ключа, NodeID, ephemeral key, client random, nonce, "
                                                      "шифртекста, повтор кадра", {})],
        [("pcap: ", {"bold": True}), ("маркер сообщения не найден; контроль прибора — служебные поля видны без TLS", {})]],
        size=14, space_after=8)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Каналы защищены TLS 1.3 со взаимной аутентификацией. Сертификаты самоподписаны, но принимаются только "
        "если ключ даёт ожидаемый NodeID. Открытые HELLO-кадры подтверждаются хэшем внутри TLS — подмену посредником "
        "видно. Сообщение дополнительно шифруется сессией отправитель–получатель поверх туннеля, поэтому "
        "ретрансляторы видят только шифртекст. Это проверено отрицательными тестами и контрольным pcap.")


def s_tunnel(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Туннель и восстановление после отказа",
          "Маршрут подтверждён подписями хопов; отказ ретранслятора — новый туннель без него")
    y = Inches(2.05)
    xs = [M + Inches(0.3) + i * Inches(1.85) for i in range(4)]
    for i, l in enumerate(["I", "R1", "R2", "D"]):
        node(s, xs[i], y, Inches(0.75), l, fill=AMBER if i == 2 else INK, size=16)
        if i < 3:
            arrow(s, xs[i] + Inches(0.8), y + Inches(0.375), xs[i + 1] - Inches(0.05), y + Inches(0.375))
    text(s, xs[2] - Inches(0.2), y - Inches(0.42), Inches(1.15), Inches(0.35), "отказ", size=14, bold=True,
         color=AMBER, align=PP_ALIGN.CENTER)
    node(s, xs[1] + Inches(0.9), y + Inches(1.35), Inches(0.75), "R3", fill=TEAL, size=16)
    text(s, xs[1] + Inches(1.75), y + Inches(1.9), Inches(2.2), Inches(0.35), "новый маршрут", size=14, bold=True,
         color=TEAL)
    arrow(s, xs[1] + Inches(0.6), y + Inches(0.8), xs[1] + Inches(0.95), y + Inches(1.45), color=TEAL)
    arrow(s, xs[1] + Inches(1.7), y + Inches(1.6), xs[3] + Inches(0.05), y + Inches(0.65), color=TEAL)
    text(s, M, Inches(4.55), Inches(6.5), Inches(1.6), [
        "BUILD по цепочке → OK с подписанными подтверждениями",
        "ACK каждой ячейки, зонд раз в секунду при простое",
        "ACTIVE → DEGRADED → DEAD → новый маршрут без отказавшего"], size=15, space_after=6)
    rf = R.STATS
    stat(s, Inches(7.5), Inches(2.0), Inches(2.6),
         f"{R.num(rf[('star', 'Отказ ретранслятора (kill): восстановление доставки, мс')]['median'], 0)} мс",
         "kill: доставка восстановлена\n(медиана star, 5 инъекций)", color=TEAL)
    stat(s, Inches(10.2), Inches(2.0), Inches(2.6),
         f"{R.num(float(rf[('star', 'Зависание ретранслятора: восстановление доставки, мс')]['median']) / 1000, 1)} с",
         "freeze: доставка восстановлена\n(медиана star, 5 инъекций)", color=AMBER)
    box(s, Inches(7.5), Inches(4.0), Inches(5.23), Inches(2.35), fill=MIST)
    text(s, Inches(7.75), Inches(4.15), Inches(4.8), Inches(2.1), [
        [("kill: ", {"bold": True}), ("ОС закрывает сокеты, сосед сразу шлёт TUNNEL_CLOSE", {})],
        [("freeze: ", {"bold": True}), ("соединения живы, видно только отсутствие ACK: 2 с тайм-аут + 1 с зонд ≈ 3,5 с", {})],
        [("повтор ", {"bold": True}), ("с тем же message_id — получатель отвечает duplicate", {})]],
        size=14, space_after=6)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Туннель строится по цепочке: каждый ретранслятор проверяет запрос и добавляет подписанное подтверждение, "
        "поэтому инициатор знает фактический маршрут. Если ретранслятор упал, сосед сразу сообщает об этом, и "
        "доставка восстанавливается за десятки миллисекунд. Если ретранслятор завис, отказ видно только по "
        "отсутствию подтверждений — около трёх с половиной секунд, это прямо следует из тайм-аутов.")


def s_plan(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Стенд и план эксперимента", "Фактор — схема bootstrap; остальное фиксировано")
    graphs = [("Underlay", "15 процессов на 127.0.0.1, полная достижимость"),
              ("Bootstrap-граф", "star: все знают узел 1; ring: узел i знает i − 1"),
              ("DHT-overlay", "содержимое k-bucket после сходимости"),
              ("Туннель", "I → R1 → R2 → D, подписи хопов")]
    y = Inches(1.95)
    for i, (h_, d) in enumerate(graphs):
        node(s, M, y, Inches(0.5), str(i + 1), fill=TEAL if i in (1, 2) else INK)
        text(s, M + Inches(0.7), y - Inches(0.02), Inches(5.5), Inches(0.35), h_, size=17, bold=True)
        text(s, M + Inches(0.7), y + Inches(0.33), Inches(5.5), Inches(0.5), d, size=14, color=MUTED)
        y += Inches(1.05)
    rows = [("5 + 5", "холодных запусков star и ring"), ("300", "контрольных lookup"),
            ("20", "инъекций отказа ретранслятора"), ("60", "сообщений по сессии")]
    for i, (v, l) in enumerate(rows):
        x = Inches(7.2) + (i % 2) * Inches(2.85)
        yy = Inches(1.95) + (i // 2) * Inches(1.6)
        stat(s, x, yy, Inches(2.6), v, l, color=INK if i % 3 == 0 else TEAL, vsize=40)
    text(s, Inches(7.2), Inches(5.3), Inches(5.5), Inches(1.0),
         "N = 15, k = 3, ALPHA = 3, R = 3, TLS 1.3, одинаковые тайм-ауты; n, min, max, медиана, среднее; p95 — "
         "только при n ≥ 20", size=13, color=MUTED)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Я различаю четыре графа: физическую достижимость — она полная, все процессы на одном компьютере; "
        "bootstrap-граф — кому какой адрес выдан; DHT-overlay — что реально лежит в таблицах; и маршрут туннеля. "
        "Звезда и кольцо — это схемы начального подключения, а не топологии. По каждой схеме пять холодных запусков, "
        "300 поисков, по десять отказов ретранслятора на схему.")


def s_results(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Результаты", f"Итоговая серия, версия кода {R.PLAN['code_version']}: сырые CSV/JSON → analyze.py")
    lk = sum(len(R.series_rows(sc, "lookups.csv")) for sc in R.PLAN["schemes"])
    lk_ok = sum(x["success"] == "True" for sc in R.PLAN["schemes"] for x in R.series_rows(sc, "lookups.csv"))
    ms = [x for sc in R.PLAN["schemes"] for x in R.series_rows(sc, "messages.csv")]
    ms_ok = sum(x["status"] == "delivered" for x in ms)
    st = R.STATS
    stats = [(f"{lk_ok}/{lk}", "lookup успешны,\nцели не было в таблице"),
             (f"{ms_ok}/{len(ms)}", "сообщений доставлено\nчерез 2 ретранслятора"),
             (f"{R.num(st[('star', 'Lookup: RPC')]['median'], 0)} RPC", "медиана на lookup,\n2 итерации"),
             (f"{R.num(st[('star', 'Сообщение по установленной сессии, мс')]['median'], 1)} мс",
              "сообщение по сессии\n(первое — ≈15 мс)")]
    for i, (v, l) in enumerate(stats):
        stat(s, M, Inches(1.85) + i * Inches(1.2), Inches(4.5), v, l, color=TEAL if i % 2 else INK, vsize=30)
    cd = CategoryChartData()
    hs = {sc: rpc_hist(sc) for sc in R.PLAN["schemes"]}
    cats = sorted({k for h, _ in hs.values() for k in h})
    cd.categories = [str(c) for c in cats]
    for sc in R.PLAN["schemes"]:
        cd.add_series(f"{sc} (n={hs[sc][1]})", [round(hs[sc][0].get(c, 0), 1) for c in cats])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.6), Inches(1.85), Inches(7.13), Inches(4.4), cd)
    ch = gf.chart
    style_chart(ch, [INK, TEAL])
    ch.has_title = True
    ch.chart_title.text_frame.text = "Число RPC на lookup, % поисков"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    pl = ch.plots[0]
    pl.has_data_labels = True
    pl.data_labels.font.size = Pt(10)
    pl.data_labels.number_format = '0'
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Все 300 контрольных поисков нашли узел, которого не было в таблице инициатора, обычно за 5 запросов и "
        "2 итерации — это согласуется с логарифмической оценкой. Все 80 сообщений доставлены через два "
        "ретранслятора. Сообщение по готовой сессии идёт около полутора миллисекунд, первое — около 15, то есть "
        "основная стоимость — установление сессии. Схема bootstrap на стоимость поиска не повлияла.")


def s_validity(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title(s, "Обсуждение и угрозы валидности", "Что объясняет числа и где их нельзя переносить")
    cards = [("≈ 2 с на отказ узла DHT", "Windows отклоняет подключение к закрытому порту через 2024–2043 мс "
              "(отдельный замер); алгоритм после этого сразу идёт к другим кандидатам", AMBER),
             ("Один компьютер", "нет сетевых задержек и потерь, процессы делят процессор: выбросы 119–179 мс — "
              "сразу после перезапуска узла", INK),
             ("Сходимость ≈ 9,8 с", "задаётся последовательным запуском процессов; star/run-04 = 14 с — два процесса "
              "стартовали по 1,9 с вместо 0,2 с", INK),
             ("Вне модели угроз", "Sybil, Eclipse, глобальный наблюдатель, анонимность — не утверждаются", TEAL)]
    for i, (h_, d, col) in enumerate(cards):
        x = M + (i % 2) * Inches(6.15)
        y = Inches(1.95) + (i // 2) * Inches(2.25)
        box(s, x, y, Inches(5.9), Inches(2.0), fill=MIST)
        node(s, x + Inches(0.25), y + Inches(0.3), Inches(0.45), "!" if col == AMBER else str(i + 1), fill=col,
             size=13)
        text(s, x + Inches(0.9), y + Inches(0.3), Inches(4.8), Inches(0.4), h_, size=17, bold=True)
        text(s, x + Inches(0.9), y + Inches(0.8), Inches(4.8), Inches(1.1), d, size=14, color=MUTED)
    footer(s, n)
    s.notes_slide.notes_text_frame.text = (
        "Два секунды при отказе узла DHT — это стек Windows: я отдельно замерил отклонение подключения к закрытому "
        "порту, оно занимает те же две секунды. Все выбросы я разобрал по журналам: это запуск процессов и общий "
        "процессор, а не алгоритм. Модель угроз сознательно не включает Sybil и Eclipse — это продвинутый уровень.")


def s_end(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, INK)
    text(s, M, Inches(0.6), W - 2 * M, Inches(0.8), "Итоги", size=36, bold=True, font=HEAD, color=PAPER)
    total, _, _ = R.junit()
    items = [("14/14", "приёмочных проверок"), (str(total), "автотестов"), ("300/300", "поисков через DHT"),
             ("80/80", "сообщений через туннель")]
    for i, (v, l) in enumerate(items):
        x = M + i * Inches(3.05)
        text(s, x, Inches(1.7), Inches(2.9), Inches(0.8), v, size=40, bold=True, font=HEAD,
             color=AMBER if i == 0 else PAPER)
        text(s, x, Inches(2.5), Inches(2.9), Inches(0.5), l, size=14, color=RGBColor(0xB8, 0xD8, 0xD3))
    text(s, M, Inches(3.5), Inches(7.5), Inches(2.3), [
        "Kademlia без центрального каталога, DHT не вырождается — прогноз 8,97 контакта подтверждён",
        "TLS 1.3 mTLS с привязкой к NodeID и E2E-защита: открытого текста в pcap нет, повтор не исполняется",
        "Отказ ретранслятора — автоматический новый маршрут; эксперимент нашёл и помог исправить дефект DHT"],
        size=16, color=PAPER, space_after=10)
    box(s, Inches(8.6), Inches(3.5), Inches(4.13), Inches(2.3), fill=RGBColor(0x1C, 0x36, 0x50))
    text(s, Inches(8.85), Inches(3.7), Inches(3.7), Inches(2.0), [
        [("Воспроизведение одной командой", {"bold": True, "color": PAPER})],
        [("python scripts/run_acceptance.py", {"font": "Courier New", "size": 13, "color": AMBER})],
        [("сеть из 15 узлов, все сценарии, сырые данные в новом каталоге", {"color": RGBColor(0xB8, 0xD8, 0xD3)})]],
        size=14, space_after=8)
    footer(s, n, dark=True)
    s.notes_slide.notes_text_frame.text = (
        "Итог: все 14 приёмочных проверок и все автотесты пройдены, сеть не использует каталог и не вырождается, "
        "данные защищены от перехвата и повтора, отказ ретранслятора обрабатывается автоматически. Всё "
        "воспроизводится одной командой. Спасибо, готов ответить на вопросы.")


def main() -> int:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    core = prs.core_properties
    core.title = "Защищённая оверлейная P2P-сеть"
    core.author = "Вихарев Д.А."
    builders = [s_title, s_goal, s_arch, s_frame, s_kademlia, s_records, s_security, s_tunnel, s_plan, s_results,
                s_validity, s_end]
    for i, b in enumerate(builders, 1):
        b(prs, i)
    prs.save(OUT)
    print(f"{OUT}: слайдов {len(prs.slides)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
