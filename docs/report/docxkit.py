"""Небольшая обёртка над python-docx для пояснительной записки: поля и шрифты по ГОСТ 7.32,
нумерация страниц, оглавление, подписи таблиц и рисунков."""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
MONO = "Consolas"


def _set_font(style_or_run, name: str, size: float | None = None, bold: bool | None = None):
    f = style_or_run.font
    f.name = name
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, "get_or_add_rPr") else None
    if rpr is not None:
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(attr), name)


def _field(paragraph, instr: str, placeholder: str = ""):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run._r.append(b)
    run._r.append(t)
    run._r.append(s)
    paragraph.add_run(placeholder)
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(e)


class Report:
    def __init__(self):
        self.doc = Document()
        self.tables = 0
        self.figures = 0
        self._setup()

    # ------------------------------------------------------------------ оформление
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.different_first_page_header_footer = True
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(p, "PAGE", "1")

        normal = d.styles["Normal"]
        _set_font(normal, FONT, 14)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.first_line_indent = Cm(1.25)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for name, size, before, after in (("Heading 1", 16, 0, 12), ("Heading 2", 14, 12, 6),
                                          ("Heading 3", 14, 6, 3)):
            st = d.styles[name]
            _set_font(st, FONT, size, True)
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.font.italic = False
            st.paragraph_format.first_line_indent = Cm(1.25)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

        settings = d.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)

    # ------------------------------------------------------------------ блоки
    def para(self, text: str = "", *, align=None, indent: bool = True, bold: bool = False, size: float | None = None,
             italic: bool = False, space_after: float = 0, keep_next: bool = False):
        p = self.doc.add_paragraph()
        self._runs(p, text, bold=bold, size=size, italic=italic)
        if align is not None:
            p.alignment = align
        if not indent:
            p.paragraph_format.first_line_indent = Cm(0)
        if space_after:
            p.paragraph_format.space_after = Pt(space_after)
        if keep_next:
            p.paragraph_format.keep_with_next = True
        return p

    def _runs(self, p, text: str, bold=False, size=None, italic=False):
        """Разметка внутри текста: **жирный**, `моноширинный`."""
        import re
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`"):
                r = p.add_run(part[1:-1])
                _set_font(r, MONO)
                r.font.size = Pt((size or 14) - 2)
            else:
                r = p.add_run(part)
                r.bold = bold or None
            if italic:
                r.italic = True
            if size and not part.startswith("`"):
                r.font.size = Pt(size)

    def h1(self, text: str, new_page: bool = True, numbered: bool = True):
        p = self.doc.add_heading(text, level=1)
        if new_page:
            p.paragraph_format.page_break_before = True       # без пустого абзаца-разрыва
        if not numbered:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
        return p

    def h2(self, text: str):
        return self.doc.add_heading(text, level=2)

    def h3(self, text: str):
        return self.doc.add_heading(text, level=3)

    def bullets(self, items: list[str], numbered: bool = False):
        for i, it in enumerate(items, 1):
            mark = f"{i}) " if numbered else "– "
            p = self.para(mark + it)
            p.paragraph_format.first_line_indent = Cm(1.25)

    def code(self, text: str, size: float = 10):
        for line in text.strip("\n").split("\n"):
            p = self.doc.add_paragraph()
            r = p.add_run(line if line else " ")
            _set_font(r, MONO, size)
            pf = p.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0.5)
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.keep_together = True
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

    def table(self, caption: str, header: list[str], rows: list[list], widths_cm: list[float] | None = None,
              size: float = 11, align_right_from: int | None = None):
        self.tables += 1
        cap = self.para(f"Таблица {self.tables} – {caption}", indent=False, keep_next=True)
        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            self._cell(t.rows[0].cells[i], h, size, bold=True, center=True)
        self._repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                right = align_right_from is not None and i >= align_right_from
                self._cell(cells[i], str(v), size, right=right)
        if widths_cm:
            t.autofit = False
            for i, w in enumerate(widths_cm):
                t.columns[i].width = Cm(w)
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
            tblw = t._tbl.tblPr.find(qn("w:tblW"))
            if tblw is None:
                tblw = OxmlElement("w:tblW")
                t._tbl.tblPr.append(tblw)
            tblw.set(qn("w:type"), "dxa")
            tblw.set(qn("w:w"), str(int(sum(widths_cm) / 2.54 * 1440)))
        self.para("", indent=False)
        return t

    def _cell(self, cell, text, size, bold=False, center=False, right=False):
        cell.text = ""
        p = cell.paragraphs[0]
        self._runs(p, text, bold=bold, size=size)
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.alignment = (WD_ALIGN_PARAGRAPH.CENTER if center else
                        WD_ALIGN_PARAGRAPH.RIGHT if right else WD_ALIGN_PARAGRAPH.LEFT)

    @staticmethod
    def _repeat_header(row):
        trpr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)

    def figure(self, path: str, caption: str, width_cm: float = 16.0):
        self.figures += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Cm(width_cm))
        c = self.para(f"Рисунок {self.figures} – {caption}", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER,
                      space_after=6)
        return c

    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def toc(self):
        p = self.para("СОДЕРЖАНИЕ", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=12)
        p.paragraph_format.page_break_before = True
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        _field(p, 'TOC \\o "1-2" \\h \\z \\u',
               "Оглавление формируется при открытии документа (если Word спросит — «Да»; иначе F9).")

    def save(self, path: str):
        self.doc.save(path)


__all__ = ["Report", "WD_ALIGN_PARAGRAPH", "Cm", "Pt", "WD_SECTION"]
